"""Выгрузка заявок в Excel.

Не CSV: владелец будет считать в этом файле суммы и строить сводные, поэтому
деньги, площади и количества уходят числами, а не текстом — иначе Excel
превратит «2 281 848 ₸» в строку, и сложить колонку не получится.

Два листа. «Заявки» — вся таблица как в админке. «Менеджеры» — тот же срез, что
на странице: кто сколько взял, продал и на какую сумму.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.db.models import Lead
from app.db.models.enums import LeadStatus

STATUS_LABEL = {
    LeadStatus.NEW: "Новая",
    LeadStatus.IN_WORK: "В работе",
    LeadStatus.WAITING: "В ожидании",
    LeadStatus.SALE: "Продажа",
    LeadStatus.REJECT: "Отказ",
    LeadStatus.NOT_TARGET: "Не целевой",
}

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

# (заголовок, ширина колонки)
_COLUMNS: list[tuple[str, int]] = [
    ("№", 6),
    ("Дата", 17),
    ("Клиент", 22),
    ("Телефон", 16),
    ("Канал / откуда", 26),
    ("UTM / кампания", 34),
    ("Google click id", 26),
    ("Город", 14),
    ("Что интересует", 26),
    ("Объём", 14),
    ("Площадь, м²", 12),
    ("Кирпича, шт", 13),
    ("Расчёт, ₸", 14),
    ("Запрос клиента", 40),
    ("Что сделано (история)", 40),
    ("Менеджер", 20),
    ("Реакция, мин", 13),
    ("Статус", 13),
    ("Продажа, ₸", 14),
    ("Причина", 32),
    ("Закрыта", 17),
]


def _local(value: datetime | None, tz_hours: int) -> datetime | None:
    """Время в рабочем часовом поясе и без tzinfo — Excel не понимает смещения."""
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(timezone(timedelta(hours=tz_hours))).replace(tzinfo=None)


def _row(lead: Lead, tz_hours: int) -> list:
    user = lead.user
    reaction = lead.reaction_seconds
    return [
        lead.id,
        _local(lead.created_at, tz_hours),
        (user.full_name if user else "") or "",
        (getattr(user, "display_phone", "") or "") if user else "",
        lead.source or "",
        lead.utm or "",
        lead.gclid or "",
        lead.city or "",
        lead.product or "",
        lead.volume or "",
        lead.area_m2,
        lead.qty_pcs,
        lead.calc_total_kzt,
        lead.request or "",
        lead.work_note or "",
        lead.taken_by_name or "",
        round(reaction / 60) if reaction is not None else None,
        STATUS_LABEL.get(lead.status, lead.status.value),
        lead.amount_kzt,
        lead.reject_reason or "",
        _local(lead.closed_at, tz_hours),
    ]


def _style_header(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"


def build_workbook(leads: list[Lead], sellers: list[dict], *, tz_hours: int = 5) -> bytes:
    """Собрать книгу и вернуть её байтами (на диск ничего не пишем)."""
    book = Workbook()

    sheet = book.active
    sheet.title = "Заявки"
    sheet.append([title for title, _ in _COLUMNS])
    _style_header(sheet, [width for _, width in _COLUMNS])

    for lead in leads:
        sheet.append(_row(lead, tz_hours))

    last = sheet.max_row
    if last > 1:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(_COLUMNS))}{last}"
    money_or_count = {12, 13, 19}
    for row in sheet.iter_rows(min_row=2, max_row=last):
        for cell in row:
            if cell.column in (2, 21):
                cell.number_format = "DD.MM.YYYY HH:MM"
            elif cell.column in money_or_count:
                cell.number_format = "# ##0"
            elif cell.column == 11:
                cell.number_format = "0.0"
        row[13].alignment = Alignment(wrap_text=True, vertical="top")

    team = book.create_sheet("Менеджеры")
    headers = [
        ("Менеджер", 24),
        ("Взял", 8),
        ("Продал", 9),
        ("Выручка, ₸", 15),
        ("Отказов", 10),
        ("Конверсия, %", 13),
        ("Реакция, мин", 13),
    ]
    team.append([title for title, _ in headers])
    _style_header(team, [width for _, width in headers])
    for seller in sellers:
        avg = seller.get("avg_reaction")
        team.append(
            [
                seller.get("name", ""),
                seller.get("taken", 0),
                seller.get("sales", 0),
                seller.get("amount", 0),
                seller.get("rejects", 0),
                seller.get("conversion", 0),
                round(avg / 60) if avg is not None else None,
            ]
        )
    for row in team.iter_rows(min_row=2, max_row=team.max_row):
        row[3].number_format = "# ##0"

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def filename(date_from: str, date_to: str) -> str:
    """Имя файла с периодом — чтобы выгрузки не путались в папке «Загрузки»."""
    return f"Заявки_{date_from}_{date_to}.xlsx"


# Имя конверсии должно совпадать с действием, заведённым в Google Ads вручную,
# иначе импорт отклонит строки. Меняется в настройках, если действие назовут иначе.
CONVERSION_NAME = "Продажа"


def google_ads_conversions_csv(
    leads: list[Lead], *, conversion_name: str = CONVERSION_NAME, tz_hours: int = 5
) -> str:
    """Офлайн-конверсии для загрузки в Google Ads: клик → живая продажа.

    Google Ads сам по себе видит только переход на сайт. Продажа случается
    через несколько дней, в переписке или по телефону, и кабинет о ней не знает —
    поэтому кампании оцениваются по кликам, а не по деньгам. Этот файл замыкает
    круг: в нём идентификатор клика (gclid), время закрытия сделки и её сумма.
    Загружается в Google Ads → Инструменты → Конверсии → Загрузки.

    В выгрузку идут только продажи с известным gclid: остальные Google не с чем
    связать, а строки без клика импорт всё равно отклонит.
    """
    header = (
        f"Parameters:TimeZone=+{tz_hours:02d}00\n"
        "Google Click ID,Conversion Name,Conversion Time,Conversion Value,Conversion Currency\n"
    )

    rows = []
    for lead in leads:
        if lead.status != LeadStatus.SALE or not lead.gclid:
            continue
        moment = _local(lead.closed_at or lead.created_at, tz_hours)
        if moment is None:
            continue
        rows.append(
            "{},{},{},{},KZT".format(
                lead.gclid,
                conversion_name,
                moment.strftime("%Y-%m-%d %H:%M:%S"),
                lead.amount_kzt or 0,
            )
        )
    return header + "\n".join(rows) + ("\n" if rows else "")
