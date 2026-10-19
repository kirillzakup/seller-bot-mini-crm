"""Выгрузка аналитического отчёта в Word (.docx).

``python-docx`` импортируется лениво: если библиотеку не поставили на сервере,
падать должна кнопка выгрузки с понятным текстом, а не весь раздел аналитики.
"""

from __future__ import annotations

import io
from typing import Any

from app.db.models import ChatReport

_METRIC_LABELS = [
    ("dialogs", "Диалогов за период"),
    ("messages", "Сообщений всего"),
    ("messages_client", "— из них от клиентов"),
    ("messages_bot", "— от бота"),
    ("messages_operator", "— от оператора"),
    ("new_leads", "Новых лидов"),
    ("with_phone", "Оставили телефон"),
    ("orders", "Заказов создано"),
    ("orders_paid", "Оплачено"),
    ("revenue_kzt", "Выручка, ₸"),
    ("conversion_pct", "Конверсия диалог → оплата, %"),
    ("unanswered_dialogs", "Диалогов без ответа"),
    ("unanswered_pct", "— доля, %"),
    ("avg_reply_seconds", "Среднее время ответа, сек"),
]

_SEVERITY = {"high": "Критично", "medium": "Важно", "low": "Мелочь"}


class DocxUnavailable(RuntimeError):
    """python-docx не установлен на сервере."""


def build_report_docx(report: ChatReport, *, shop_name: str = "") -> bytes:
    """Собирает .docx с метриками, находками и предложениями. Возвращает байты файла."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as e:  # pragma: no cover - зависит от окружения
        raise DocxUnavailable(
            "Для выгрузки в Word нужна библиотека python-docx. "
            "Установите её на сервере: pip install python-docx"
        ) from e

    doc = Document()

    title = doc.add_heading("Аналитический отчёт по перепискам", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    sub = doc.add_paragraph()
    sub.add_run(
        f"{shop_name + ' · ' if shop_name else ''}"
        f"Период: {report.period_from:%d.%m.%Y} — {report.period_to:%d.%m.%Y}\n"
    ).bold = True
    sub.add_run(
        f"Разобрано диалогов: {report.dialogs_analyzed} · "
        f"сообщений в периоде: {report.messages_analyzed} · "
        f"отчёт сформирован {report.created_at:%d.%m.%Y %H:%M}"
    ).font.size = Pt(9)

    # --- Метрики ---
    doc.add_heading("Показатели за период", level=1)
    metrics: dict[str, Any] = report.metrics or {}
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for key, label in _METRIC_LABELS:
        if key not in metrics:
            continue
        row = table.add_row().cells
        row[0].text = label
        value = metrics.get(key)
        row[1].text = f"{value:,}".replace(",", " ") if isinstance(value, int) else str(value)

    for key, heading in (
        ("by_channel", "По каналам"),
        ("by_source", "По источникам"),
        ("by_stage", "По стадиям воронки"),
    ):
        data = metrics.get(key) or {}
        if isinstance(data, dict) and data:
            doc.add_paragraph(
                f"{heading}: " + ", ".join(f"{k} — {v}" for k, v in data.items()),
                style="List Bullet",
            )

    # --- Резюме ---
    if report.summary:
        doc.add_heading("Главное", level=1)
        doc.add_paragraph(report.summary)

    # --- Находки ---
    findings = report.findings or []
    if findings:
        doc.add_heading("Что мешает продавать", level=1)
        for i, f in enumerate(findings, 1):
            if not isinstance(f, dict):
                continue
            sev = _SEVERITY.get(str(f.get("severity", "")).lower(), "")
            head = doc.add_paragraph()
            head.add_run(f"{i}. {f.get('title', 'Без названия')}").bold = True
            if sev:
                head.add_run(f"  [{sev}]").italic = True
            if f.get("evidence"):
                q = doc.add_paragraph(f"«{f['evidence']}»", style="Intense Quote")
                q.paragraph_format.space_after = Pt(4)
            if f.get("recommendation"):
                doc.add_paragraph(f"Что делать: {f['recommendation']}", style="List Bullet")

    # --- Предложения ---
    proposals = report.proposals or {}
    script = (proposals.get("sales_script") or "").strip()
    settings = proposals.get("settings") or {}
    if script or settings:
        doc.add_heading("Предложенные изменения", level=1)
        if settings:
            doc.add_paragraph("Настройки:").runs[0].bold = True
            for key, value in settings.items():
                doc.add_paragraph(f"{key} → {value}", style="List Bullet")
        if script:
            doc.add_paragraph("Новый скрипт продаж:").runs[0].bold = True
            for line in script.splitlines():
                doc.add_paragraph(line)

    segments = proposals.get("segments") or []
    if segments:
        doc.add_heading("Сегменты для ретаргета", level=1)
        for s in segments:
            if not isinstance(s, dict):
                continue
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(str(s.get("name", ""))).bold = True
            if s.get("criteria"):
                p.add_run(f" — {s['criteria']}")
            ids = s.get("user_ids") or []
            if isinstance(ids, list) and ids:
                p.add_run(f" (контактов: {len(ids)})").italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
