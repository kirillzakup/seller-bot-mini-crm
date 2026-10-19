"""Заявки продавцам: когда заводятся, что в карточке, как считается сводка."""

from __future__ import annotations

import asyncio
from datetime import UTC, datetime, timedelta

from app.admin.routes.leads import _seller_rows
from app.db.models.enums import Channel, LeadStatus
from app.services import leads_service as leads


def _make_user(session, **kw):
    from app.db.models import User

    user = User(channel=Channel.WHATSAPP, external_id=kw.pop("ext", "77001234567"), **kw)
    session.add(user)
    return user


def test_handoff_opens_one_lead_and_card_has_contacts(client, monkeypatch):
    """Пока заявка открыта, повторная передача менеджеру не плодит вторую."""
    from app.db.base import AsyncSessionLocal

    posted: list = []

    async def fake_post(session, lead, user):
        posted.append(lead.id)
        lead.card_message_id = 1000 + len(posted)
        return True

    monkeypatch.setattr(leads, "post_card", fake_post)

    async def run():
        async with AsyncSessionLocal() as session:
            user = _make_user(session, first_name="Пётр", ext="77071112233")
            await session.flush()

            first = await leads.record_handoff(session, user, reason="Просит оператора")
            assert first is not None and first.status == LeadStatus.NEW

            again = await leads.record_handoff(session, user, reason="Просит перезвонить")
            assert again.id == first.id, "та же заявка, а не вторая карточка"
            assert "Просит перезвонить" in again.request, "причина дописалась в ту же заявку"

            first.status = LeadStatus.SALE  # закрыли — клиент вернулся позже
            await session.flush()
            third = await leads.record_handoff(session, user, reason="Снова просит расчёт")
            assert third.id != first.id, "после закрытия открывается новая"

            text = leads.card_text(first, user)
            assert "Заявка #" in text
            assert "77071112233" in text
            assert "wa.me/77071112233" in text
            await session.rollback()

    asyncio.run(run())
    assert len(posted) == 2, "карточка на каждую незакрытую заявку, а не на каждую передачу"


def test_reaction_time_and_seller_summary():
    """Сводка по менеджерам считает только то, что человек реально забрал."""
    from app.db.models import Lead

    now = datetime.now(UTC)

    def lead(**kw):
        row = Lead(user_id=1, **kw)
        row.created_at = kw.pop("created", now - timedelta(minutes=10))
        return row

    sold = lead(status=LeadStatus.SALE, amount_kzt=1_450_000, taken_by_name="Асем")
    sold.taken_at = sold.created_at + timedelta(minutes=4)
    rejected = lead(status=LeadStatus.REJECT, taken_by_name="Асем", reject_reason="дорого")
    rejected.taken_at = rejected.created_at + timedelta(minutes=6)
    nobody = lead(status=LeadStatus.NEW)

    assert sold.reaction_seconds == 240
    assert nobody.reaction_seconds is None

    # Из SQLite created_at приходит без таймзоны, а обработчик кнопки ставит
    # datetime.now(UTC): раньше это роняло card_text, и карточка не обновлялась.
    mixed = lead(status=LeadStatus.IN_WORK, taken_by_name="Асем")
    mixed.created_at = mixed.created_at.replace(tzinfo=None)
    mixed.taken_at = datetime.now(UTC)
    assert mixed.reaction_seconds is not None
    assert "Асем" in leads.card_text(mixed, None)
    assert leads._days_in_work(mixed) == ""
    assert sold.has_result and not nobody.has_result

    rows = _seller_rows([sold, rejected, nobody])
    assert len(rows) == 1, "ничья заявка не попадает ни к кому в статистику"
    row = rows[0]
    assert row["name"] == "Асем"
    assert (row["taken"], row["sales"], row["rejects"]) == (2, 1, 1)
    assert row["amount"] == 1_450_000
    assert row["conversion"] == 50
    assert row["avg_reaction"] == 300


def test_admin_page_renders(client):
    client.post("/admin/login", data={"username": "admin", "password": "admin"})
    r = client.get("/admin/leads")
    assert r.status_code == 200
    assert "Заявки" in r.text
    assert "Конверсия" in r.text


class _TgUser:
    """Минимум от aiogram-пользователя, который нужен реестру менеджеров."""

    def __init__(self, uid, full_name="", username=None):
        self.id = uid
        self.full_name = full_name
        self.username = username


def test_seller_registry_replaces_telegram_nickname(client):
    """Пока менеджер не представился, в отчёт идёт его ник — и это видно по флагу."""
    from app.db.base import AsyncSessionLocal

    async def run():
        async with AsyncSessionLocal() as session:
            raw = _TgUser(777, full_name="🔥ASEM🔥")

            name, registered = await leads.seller_name(session, raw)
            assert (name, registered) == ("🔥ASEM🔥", False)

            await leads.register_seller(session, 777, "Асем Нурланова")
            name, registered = await leads.seller_name(session, raw)
            assert (name, registered) == ("Асем Нурланова", True)

            assert await leads.sellers(session) == {"777": "Асем Нурланова"}

            assert await leads.remove_seller(session, 777) is True
            assert await leads.remove_seller(session, 777) is False
            name, registered = await leads.seller_name(session, raw)
            assert (name, registered) == ("🔥ASEM🔥", False)
            await session.rollback()

    asyncio.run(run())


def test_seller_name_falls_back_when_profile_is_empty(client):
    from app.db.base import AsyncSessionLocal

    async def run():
        async with AsyncSessionLocal() as session:
            assert await leads.seller_name(session, None) == ("менеджер", False)
            nameless = _TgUser(9, full_name="", username="asem_kz")
            assert await leads.seller_name(session, nameless) == ("asem_kz", False)

    asyncio.run(run())


# ── Расчёт в переписке ────────────────────────────────────────────────────


def test_every_handoff_becomes_a_card(client, monkeypatch):
    """Оператор, звонок, вопрос без ответа — всё это одна воронка с кнопками."""
    from app.db.base import AsyncSessionLocal
    from app.sales.tools import run_tool

    cards: list = []

    async def fake_post(session, lead, user):
        cards.append(lead.id)
        lead.card_message_id = 500 + len(cards)
        return True

    async def no_op_message(session, text):
        return False  # чат уведомлений не настроен — проверяем именно карточку

    monkeypatch.setattr(leads, "post_card", fake_post)
    monkeypatch.setattr("app.sales.tools.send_operator_message", no_op_message)

    async def run():
        async with AsyncSessionLocal() as session:
            for ext, tool, args in (
                ("77011110001", "escalate_to_operator", {"reason": "Просит человека"}),
                ("77011110002", "request_callback", {"note": "Перезвоните после 18"}),
                ("77011110003", "ask_admin", {"question": "Есть ли доставка в Тараз?"}),
            ):
                user = _make_user(session, first_name="Клиент", ext=ext)
                await session.flush()
                await run_tool(tool, args, session=session, user=user, messenger=None)
                lead = await leads.open_lead(session, user)
                assert lead is not None, f"{tool} не завёл заявку"
                assert lead.card_message_id, f"{tool}: карточка не отправлена"
            await session.rollback()

    asyncio.run(run())
    assert len(cards) == 3, "по карточке на каждую передачу менеджеру"


def test_bot_does_not_claim_delivery_when_card_failed(client, monkeypatch):
    """Карточка не ушла — бот не должен говорить клиенту «оператор уведомлён»."""
    from app.db.base import AsyncSessionLocal
    from app.sales.tools import run_tool

    async def failed_post(session, lead, user):
        return False  # чат продаж не настроен

    async def no_op_message(session, text):
        return False

    monkeypatch.setattr(leads, "post_card", failed_post)
    monkeypatch.setattr("app.sales.tools.send_operator_message", no_op_message)

    async def run():
        async with AsyncSessionLocal() as session:
            user = _make_user(session, first_name="Клиент", ext="77011110009")
            await session.flush()
            reply = await run_tool(
                "escalate_to_operator",
                {"reason": "Просит человека"},
                session=session,
                user=user,
                messenger=None,
            )
            assert "ВНИМАНИЕ" in reply
            assert "НЕ доставлено" in reply
            await session.rollback()

    asyncio.run(run())


def test_card_offers_four_buttons_after_takeover():
    """ТЗ: после «Забрал в работу» — четыре кнопки, включая «В работе»."""
    from app.db.models import Lead

    lead = Lead(user_id=1, status=LeadStatus.IN_WORK)
    lead.id = 7
    buttons = [b.text for row in leads.card_kb(lead).inline_keyboard for b in row]
    assert buttons == ["⏳ В работе", "💰 Продажа", "❌ Отказ", "🚫 Не целевой"]
    assert [b.callback_data for row in leads.card_kb(lead).inline_keyboard for b in row] == [
        "lead:work:7",
        "lead:sale:7",
        "lead:reject:7",
        "lead:nt:7",
    ]

    fresh = Lead(user_id=1, status=LeadStatus.NEW)
    fresh.id = 8
    assert [b.text for row in leads.card_kb(fresh).inline_keyboard for b in row] == [
        "🙋 Забрал в работу"
    ]


def test_not_target_reason_shows_up_on_the_card():
    """У «не целевого» тоже видна причина — иначе непонятно, кого отсеиваем."""
    from app.db.models import Lead

    lead = Lead(user_id=1, status=LeadStatus.NOT_TARGET, taken_by_name="Асем")
    lead.id = 9
    lead.reject_reason = "Посредник, перепродаёт"
    text = leads.card_text(lead, None)
    assert "Посредник, перепродаёт" in text
    assert "🚫 Причина" in text


def test_daily_reminder_fires_once_per_day_at_the_set_hour(client, monkeypatch):
    """Напоминание уходит в назначенный час и ровно один раз за сутки."""
    from datetime import datetime, timedelta, timezone

    from app.db.base import AsyncSessionLocal
    from app.services import jobs
    from app.services.settings_service import set_setting

    kz = timezone(timedelta(hours=5))
    clock = {"now": datetime(2026, 9, 18, 8, 30, tzinfo=kz)}  # 08:30 — ещё рано

    class FrozenDatetime(datetime):
        @classmethod
        def now(cls, tz=None):
            return clock["now"].astimezone(tz) if tz else clock["now"]

    monkeypatch.setattr(jobs, "datetime", FrozenDatetime)

    reminded: list[int] = []

    async def fake_remind(session):
        reminded.append(1)
        return 1

    monkeypatch.setattr("app.services.leads_service.remind_leads_in_work", fake_remind)

    async def run():
        # Саму рассылку подменили выше, заявки здесь не нужны — проверяем расписание.
        async with AsyncSessionLocal() as session:
            await set_setting(session, jobs.LEADS_REMIND_DATE_KEY, None)
            await session.commit()

        assert await jobs.remind_leads_in_work() == 0, "в 08:30 ещё рано"

        clock["now"] = datetime(2026, 9, 18, 9, 5, tzinfo=kz)
        assert await jobs.remind_leads_in_work() == 1, "в 09:05 пора"

        clock["now"] = datetime(2026, 9, 18, 9, 40, tzinfo=kz)
        assert await jobs.remind_leads_in_work() == 0, "второй раз за сутки не шлём"

        clock["now"] = datetime(2026, 9, 19, 9, 1, tzinfo=kz)
        assert await jobs.remind_leads_in_work() == 1, "назавтра снова напоминаем"
        assert len(reminded) == 2

    asyncio.run(run())


def test_morning_digest_is_a_mini_report_and_two_lists(client, monkeypatch):
    """В 9:00 — мини-отчёт с цифрами, список «Не забраны» и список «В работе»;
    номер заявки — ссылка на её карточку, без карточки — ссылка на переписку."""
    from datetime import UTC, datetime, timedelta

    from app.db.base import AsyncSessionLocal
    from app.db.models import Lead
    from app.services import leads_reminders, leads_service

    sent: list[dict] = []

    class FakeBot:
        async def send_message(
            self, chat_id, text, reply_markup=None, message_thread_id=None, **kw
        ):
            sent.append({"text": text, "kb": reply_markup, "thread": message_thread_id, "kw": kw})

    async def targets(session):
        return [(-100123, 7)]

    monkeypatch.setattr(leads_service, "leads_targets", targets)
    monkeypatch.setattr(leads_reminders, "_SEND_PAUSE_S", 0)

    async def run():
        async with AsyncSessionLocal() as session:
            user = _make_user(session, first_name="Ержан", ext="77015556677")
            await session.flush()
            untaken = Lead(user_id=user.id, status=LeadStatus.NEW, city="Семей")
            untaken.created_at = datetime.now(UTC) - timedelta(days=2)
            working = Lead(
                user_id=user.id,
                status=LeadStatus.IN_WORK,
                taken_by_name="Асем",
                card_chat_id=-1001111111111,
                card_message_id=555,
            )
            working.taken_at = datetime.now(UTC) - timedelta(days=3)
            parked = Lead(
                user_id=user.id,
                status=LeadStatus.WAITING,
                remind_at=datetime.now(UTC) + timedelta(days=5),
            )
            session.add_all([untaken, working, parked])
            await session.flush()

            count = await leads_service.remind_leads_in_work(session, bot=FakeBot())
            texts = [m["text"] for m in sent]
            # В общей тестовой базе бывают чужие заявки — точные цифры не сверяем.
            assert count == len(sent) >= 3
            assert all(m["thread"] == 7 for m in sent), "всё в тему заявок"
            assert all(m["kb"] is None for m in sent), "кнопки — в самих карточках"
            assert all(m["kw"]["link_preview_options"].is_disabled for m in sent)
            report = texts[0]
            assert report.startswith("☀️ <b>Заявки на утро")
            for line in ("🆕 Не забраны:", "⏳ В работе:", "🕓 В ожидании:", "Вчера:"):
                assert line in report
            new_list = next(t for t in texts if t.startswith("🆕 <b>Не забраны"))
            work_list = next(t for t in texts if t.startswith("⏳ <b>В работе"))
            assert texts.index(new_list) < texts.index(work_list), "незабранные — первыми"
            assert f"#{untaken.id} Ержан · 2 дн" in new_list
            assert f"/admin/dialogs?u={user.id}" in new_list, "без карточки — в переписку"
            assert f'<a href="https://t.me/c/1111111111/555">#{working.id}</a>' in work_list
            assert "Асем · 3 дн" in work_list
            assert f"#{parked.id}" not in new_list + work_list, "ожидание — не в списках"
            await session.rollback()

    asyncio.run(run())


def test_work_notes_keep_history_and_card_shows_the_latest():
    from datetime import UTC, datetime

    from app.bot.handlers.leads import _ASK_RE
    from app.db.models import Lead

    lead = Lead(user_id=1, status=LeadStatus.IN_WORK)
    leads.add_work_note(lead, "Ксения", "звонила, не взял", datetime(2026, 10, 7, tzinfo=UTC))
    leads.add_work_note(lead, "Ксения", "отправила расчёт", datetime(2026, 10, 8, tzinfo=UTC))
    assert lead.work_note.splitlines() == [
        "07.10 Ксения: звонила, не взял",
        "08.10 Ксения: отправила расчёт",
    ]
    assert leads.latest_work_note(lead) == "08.10 Ксения: отправила расчёт"
    card = leads.card_text(lead, None)
    assert "отправила расчёт" in card and "звонила" not in card, "в карточке — последнее"
    assert _ASK_RE.search("⟦заявка #12 сделано⟧").groups() == ("12", "сделано")


def test_excel_export_keeps_numbers_as_numbers(client, monkeypatch):
    """В выгрузке деньги и количества — числа, иначе их не сложить в Excel."""
    import io as _io
    from datetime import UTC, datetime, timedelta

    from openpyxl import load_workbook

    from app.admin.routes.leads import _seller_rows
    from app.db.base import AsyncSessionLocal
    from app.db.models import Lead
    from app.services import leads_export

    async def run():
        async with AsyncSessionLocal() as session:
            user = _make_user(session, first_name="Ержан", ext="77016667788")
            await session.flush()
            lead = Lead(
                user_id=user.id,
                status=LeadStatus.SALE,
                taken_by_name="Асем Нурланова",
                city="Павлодар",
                area_m2=112.0,
                qty_pcs=6116,
                calc_total_kzt=2495328,
                amount_kzt=2400000,
                calc_text="112 м² → 6 116 шт",
            )
            lead.taken_at = datetime.now(UTC)
            lead.created_at = datetime.now(UTC) - timedelta(minutes=8)
            session.add(lead)
            await session.flush()
            await session.refresh(lead, ["user"])
            return leads_export.build_workbook([lead], _seller_rows([lead]))

    payload = asyncio.run(run())
    book = load_workbook(_io.BytesIO(payload))

    assert book.sheetnames == ["Заявки", "Менеджеры"]
    sheet = book["Заявки"]
    headers = [c.value for c in sheet[1]]
    assert headers[:4] == ["№", "Дата", "Клиент", "Телефон"]
    assert "UTM / кампания" in headers and "Google click id" in headers

    row = {headers[i]: c.value for i, c in enumerate(sheet[2])}
    assert row["Клиент"] == "Ержан"
    assert row["Город"] == "Павлодар"
    assert row["Площадь, м²"] == 112.0, "число, а не строка «112 м²»"
    assert row["Кирпича, шт"] == 6116
    assert row["Расчёт, ₸"] == 2495328
    assert row["Продажа, ₸"] == 2400000
    assert row["Статус"] == "Продажа"
    assert isinstance(row["Реакция, мин"], int)

    team = book["Менеджеры"]
    seller = {c.value: team.cell(row=2, column=i + 1).value for i, c in enumerate(team[1])}
    assert seller["Менеджер"] == "Асем Нурланова"
    assert seller["Выручка, ₸"] == 2400000
    assert seller["Конверсия, %"] == 100


def test_excel_export_route_returns_a_file(client):
    client.post("/admin/login", data={"username": "admin", "password": "admin"})
    r = client.get("/admin/leads/export.xlsx?date_from=2026-01-01&date_to=2026-12-31")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]
    assert r.content[:2] == b"PK", "это настоящий xlsx, а не html с ошибкой"


def test_google_ads_conversions_only_include_sales_with_a_click_id():
    """В Google Ads уходят только продажи, у которых есть идентификатор клика."""
    from datetime import UTC, datetime

    from app.db.models import Lead
    from app.services.leads_export import google_ads_conversions_csv

    def make(status, amount, gclid):
        lead = Lead(user_id=1, status=status, amount_kzt=amount, gclid=gclid)
        lead.created_at = datetime(2026, 9, 17, 9, 20, tzinfo=UTC)
        lead.closed_at = datetime(2026, 9, 17, 14, 20, tzinfo=UTC)
        return lead

    rows = google_ads_conversions_csv(
        [
            make(LeadStatus.SALE, 2_400_000, "Cj0KCQjwTEST"),
            make(LeadStatus.SALE, 900_000, ""),  # клик неизвестен — Google не свяжет
            make(LeadStatus.REJECT, 0, "Cj0KCQjwOTHER"),  # не продажа
            make(LeadStatus.IN_WORK, 0, "Cj0KCQjwWORK"),  # ещё не закрыта
        ]
    ).strip().splitlines()

    assert rows[0] == "Parameters:TimeZone=+0500"
    assert rows[1].startswith("Google Click ID,Conversion Name,Conversion Time")
    assert len(rows) == 3, "ровно одна конверсия"
    click, name, moment, value, currency = rows[2].split(",")
    assert click == "Cj0KCQjwTEST"
    assert name == "Продажа"
    assert moment == "2026-09-17 19:20:00", "время закрытия сделки, в местной зоне"
    assert (value, currency) == ("2400000", "KZT")


def test_ad_tags_carry_over_to_the_next_lead(client, monkeypatch):
    """Клиент кликнул по рекламе один раз, а обратился дважды — метка не теряется."""
    from app.db.base import AsyncSessionLocal

    async def fake_post(session, lead, user):
        lead.card_message_id = 42
        return True

    monkeypatch.setattr(leads, "post_card", fake_post)

    async def run():
        async with AsyncSessionLocal() as session:
            first = await leads.record_web_lead(
                session,
                name="Ержан",
                phone="+7 701 555 44 33",
                source="Окно расчёта",
                gclid="Cj0KCQjwFIRST",
                utm="utm_source=google&utm_campaign=23827100596",
            )
            assert first.gclid == "Cj0KCQjwFIRST"

            first.status = LeadStatus.SALE  # заявку закрыли
            await session.flush()

            from app.db.models import User

            user = await session.get(User, first.user_id)
            second = await leads.record_handoff(session, user, reason="Снова просит расчёт")
            assert second.id != first.id
            assert second.gclid == "Cj0KCQjwFIRST", "реклама привела и этого клиента тоже"
            await session.rollback()

    asyncio.run(run())
